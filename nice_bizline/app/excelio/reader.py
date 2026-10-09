"""입력 엑셀 읽기 - 회사명/사업자번호/대표자명 컬럼을 유연하게 매칭."""
from __future__ import annotations

import re
import openpyxl


_NORMALIZE = re.compile(r"[\s()\[\]]")


def _norm(s: str) -> str:
    return _NORMALIZE.sub("", str(s)).lower()


_HEADER_ALIASES = {
    "회사명": {"회사명", "업체명", "상호", "회사", "기업명", "업체", "고객사", "고객사명",
              "거래처", "거래처명", "companyname", "company"},
    "사업자번호": {"사업자번호", "사업자등록번호", "사업자등록", "bizno", "businessno"},
    "대표자명": {"대표자명", "대표자", "대표", "ceo", "representative"},
    "주소": {"주소", "소재지", "본사주소", "사업장주소", "address", "addr"},
    "전화번호": {"전화번호", "전화", "대표번호", "연락처", "tel", "phone"},
}


def _match_header(cell_value) -> str | None:
    n = _norm(cell_value or "")
    if not n:
        return None
    for canonical, aliases in _HEADER_ALIASES.items():
        if any(_norm(a) == n for a in aliases):
            return canonical
    return None


# 중복(동명) 필터에 쓸 수 있는 컬럼(회사명 제외). 검색 결과에 있는 값이라 사전 필터 가능.
FILTERABLE_FIELDS = ["대표자명", "주소"]


def available_filter_fields(companies: list[dict]) -> list[str]:
    """읽어들인 목록에서 값이 하나라도 있는 '필터 가능' 컬럼을 반환.

    UI에서 이 목록을 체크박스로 보여주고, 사용자가 고른 것만 중복 필터에 쓴다.
    """
    out = []
    for f in FILTERABLE_FIELDS:
        if any((c.get(f) or "").strip() for c in companies):
            out.append(f)
    return out


def _column_map(header_row) -> dict[int, str]:
    """헤더 행 → {열 번호: 표준 이름}. 회사명 헤더가 없으면 1열을 회사명으로 본다."""
    col_map: dict[int, str] = {}
    for idx, val in enumerate(header_row):
        canon = _match_header(val)
        if canon:
            col_map[idx] = canon
    if "회사명" not in col_map.values():
        # 회사명 헤더만 인식 실패 시 첫 컬럼을 회사명으로 간주하되,
        # 이미 인식된 다른 컬럼(주소/대표자명 등)은 보존한다.
        col_map = {i: c for i, c in col_map.items() if i != 0}
        col_map[0] = "회사명"
    return col_map


def inspect_columns(path: str) -> list[dict]:
    """화면의 '이렇게 읽었습니다' 카드용: 엑셀 열마다 어떻게 인식했는지.

    반환: [{"original": 원래 열 이름, "field": 표준 이름 또는 None,
            "renamed": 원래 이름과 표준 이름이 다르면 True, "guessed": 헤더로 못 찾아 1열로 본 경우 True}]
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    try:
        header_row = next(wb.active.iter_rows(values_only=True), None)
    finally:
        wb.close()
    if not header_row:
        return []
    col_map = _column_map(header_row)
    out = []
    for idx, val in enumerate(header_row):
        original = "" if val is None else str(val).strip()
        if not original and idx not in col_map:
            continue
        field = col_map.get(idx)
        out.append({
            "original": original or f"{idx + 1}열",
            "field": field,
            "renamed": bool(field) and _norm(original) != _norm(field),
            "guessed": field == "회사명" and _match_header(val) != "회사명",
        })
    return out


def read_company_list(path: str) -> list[dict]:
    """첫 시트 1행을 헤더로 보고, 2행부터 데이터를 읽어 dict 리스트로 반환.

    회사명은 필수. 사업자번호/대표자명은 있으면 매칭 정확도 향상.
    """
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    ws = wb.active

    rows = ws.iter_rows(values_only=True)
    header_row = next(rows, None)
    if not header_row:
        wb.close()
        return []

    col_map = _column_map(header_row)

    out: list[dict] = []
    for row in rows:
        rec: dict = {}
        for idx, key in col_map.items():
            if idx < len(row) and row[idx] is not None:
                rec[key] = str(row[idx]).strip()
        if rec.get("회사명"):
            out.append(rec)
    wb.close()
    return out
